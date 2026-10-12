#!/usr/bin/env python3
"""Render satu naskah Mode A 30 Kunci ke DOCX. Pemakaian: python3 render_bab_docx.py BAB_056_....md
Standar Boss 2026-10-09/10: Latin justify; Arab rata kanan satu paragraf menyambung, run w:rtl
+Scheherazade New; nomor ayat kurung biasa (..) di BELAKANG tiap ayat (dalam Arab angka Arab-Indic).
Jangan diubah tanpa alasan — pola ini sudah lolos QC visual Bab 55."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

NB = '\u00a0'

def make_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(11)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def para(doc, text, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(); p.alignment = align
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def ar_run(p, text):
    r = p.add_run(text); r.font.name = 'Scheherazade New'; r.font.size = Pt(15)
    rPr = r._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn('w:rtl'), {})); rPr.append(rPr.makeelement(qn('w:cs'), {qn('w:val'): '1'}))
    return r

def arab_flowing(doc, text):
    """Satu paragraf menyambung, rata KANAN; nomor ayat menempel di belakang tiap ayat."""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    chunks = [x for x in re.split(r'(?<=\))\s+', text) if x]
    for k, c in enumerate(chunks):
        m = re.match(r'^(.*?)\s*(\(.+\))$', c)
        if m:
            ar_run(p, (NB + m.group(1)) if k else m.group(1))
            ar_run(p, NB + m.group(2))
        else:
            ar_run(p, (NB if k else '') + c)
    pPr = p._p.get_or_add_pPr(); pPr.append(pPr.makeelement(qn('w:rtl'), {}))

def render(md_path, out_path):
    t = open(md_path, encoding='utf-8').read()
    doc = make_doc()
    lines = t.split('\n')
    i = 0
    while i < len(lines):
        l = lines[i]
        if not l.strip():
            i += 1; continue
        if l.startswith('# '):
            para(doc, l[2:].strip(), bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif l.startswith('## '):
            para(doc, l[3:].strip(), bold=True, size=14)
        elif l.startswith('::: dalil'):
            i += 1; buf = []
            while i < len(lines) and lines[i].strip() != ':::':
                buf.append(lines[i]); i += 1
            for b in buf:
                if b.startswith('ARAB: '):
                    arab_flowing(doc, b[6:].strip())
                elif b.startswith('Artinya: '):
                    para(doc, 'Artinya: ' + b[9:].strip(), italic=True)
                elif re.match(r'^(RUJUKAN|DERAJAT|STATUS|PERAWI|LATIN): ', b):
                    para(doc, b.strip(), size=9)
                elif b.startswith('QS. '):
                    para(doc, b.strip(), size=9)
        elif l.startswith('::: '):
            nama = l[4:].strip()
            i += 1; buf = []
            while i < len(lines) and lines[i].strip() != ':::':
                buf.append(lines[i]); i += 1
            para(doc, nama.replace('-', ' ').title(), bold=True, size=12)
            for b in buf:
                if not b.strip(): continue
                if b.startswith('ARAB: '):
                    arab_flowing(doc, b[6:].strip())
                elif b.startswith('Artinya: '):
                    para(doc, 'Artinya: ' + b[9:].strip(), italic=True)
                elif b.startswith('|'):
                    tbl = []
                    while i < len(lines) and False: pass
                    # kumpulkan seluruh blok tabel berurutan dari buf
                    j = buf.index(b); rows = []
                    while j < len(buf) and buf[j].startswith('|'):
                        rows.append([c.strip() for c in buf[j].strip().strip('|').split('|')]); j += 1
                    rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
                    if rows:
                        tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = 'Table Grid'
                        for ri, row in enumerate(rows):
                            for ci, cell in enumerate(row[:len(rows[0])]):
                                cellp = tb.rows[ri].cells[ci].paragraphs[0]
                                cellp.text = cell
                                if ri == 0:
                                    for r in cellp.runs: r.bold = True
                    i = i  # noqa
                    # majukan i sesuai jumlah baris tabel yang sudah dipakai dari buf
                    i += (j - buf.index(b)) - 1
                    break
                else:
                    para(doc, b.strip())
        elif l.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i += 1
            rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
            tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row[:len(rows[0])]):
                    cellp = tb.rows[ri].cells[ci].paragraphs[0]
                    cellp.text = cell
                    if ri == 0:
                        for r in cellp.runs: r.bold = True
            continue
        elif l.startswith('*') and l.endswith('*') and len(l) > 2:
            para(doc, l.strip('*'), italic=True)
        else:
            para(doc, l.strip())
        i += 1
    doc.save(out_path)
    print('OK', out_path)

if __name__ == '__main__':
    render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1].replace('.md', '.docx'))
