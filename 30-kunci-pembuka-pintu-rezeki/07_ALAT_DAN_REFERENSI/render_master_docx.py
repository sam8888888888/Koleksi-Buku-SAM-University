#!/usr/bin/env python3
"""Render MASTER_NASKAH_v1.md (front matter + 105 bab + 34 lampiran) ke satu DOCX.
Menggabungkan kemampuan render_bab_docx (blok dalil Arab flowing) dan
render_lampiran_docx (tabel markdown). \newpage = page break."""
import re, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

NB = '\u00a0'

def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(11)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def ar_run(p, text, size=15):
    r = p.add_run(text); r.font.name = 'Scheherazade New'; r.font.size = Pt(size)
    rPr = r._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn('w:rtl'), {})); rPr.append(rPr.makeelement(qn('w:cs'), {qn('w:val'): '1'}))
    return r

def para(doc, text, bold=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(); p.alignment = align
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    return p

def arab_flowing(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    chunks = [x for x in re.split(r'(?<=\))\s+', text) if x]
    for k, c in enumerate(chunks):
        m = re.match(r'^(.*?)\s*(\(.+.\))$', c)
        if m:
            ar_run(p, (NB + m.group(1)) if k else m.group(1))
            ar_run(p, NB + m.group(2))
        else:
            ar_run(p, (NB if k else '') + c)
    pPr = p._p.get_or_add_pPr(); pPr.append(pPr.makeelement(qn('w:rtl'), {}))

def add_mixed(p, text, size=10.5):
    for seg in re.split(r'(\*\*.+?\*\*)', text):
        if not seg: continue
        if seg.startswith('**'):
            r = p.add_run(seg[2:-2]); r.bold = True; r.font.size = Pt(size)
        else:
            r = p.add_run(seg); r.font.size = Pt(size)

def render(md_path, out_path):
    t = open(md_path, encoding='utf-8').read()
    doc = make_doc()
    lines = t.split('\n')
    i = 0
    while i < len(lines):
        l = lines[i]
        if not l.strip():
            i += 1; continue
        if l.strip() == '\\newpage':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif l.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                i += 1
        elif l.startswith('# '):
            para(doc, l[2:].strip(), bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif l.startswith('## '):
            para(doc, l[3:].strip(), bold=True, size=14)
        elif l.startswith('### '):
            para(doc, l[4:].strip(), bold=True, size=12)
        elif l.startswith('::: dalil'):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith(':::'):
                buf.append(lines[i].strip()); i += 1
            block = '\n'.join(buf)
            ar = [x for x in buf if re.search(r'[\u0600-\u06FF]', x)]
            if ar:
                arab_flowing(doc, ' '.join(x for x in buf if re.search(r'[\u0600-\u06FF]', x) or x.startswith(('QS.', 'HR.', 'Artinya:', 'DERAJAT:', 'RUJUKAN:', 'STATUS:'))))
            else:
                para(doc, block.replace('\n', ' '))
        elif re.match(r'^::: \w', l):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith(':::'):
                buf.append(lines[i]); i += 1
            for b in buf:
                if b.strip(): para(doc, b.strip())
        elif l.startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].startswith('|'):
                tbl.append(lines[i]); i += 1
            rows = [x for x in tbl if not re.match(r'^\|[\s\-:|]+\|$', x)]
            if not rows: continue
            ncols = max(len(x.split('|')) - 2 for x in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, rl in enumerate(rows):
                cells = [c.strip() for c in rl.strip().strip('|').split('|')]
                tr = table.add_row()
                for j in range(ncols):
                    txt = cells[j] if j < len(cells) else ''
                    tp = tr.cells[j].paragraphs[0]
                    if re.search(r'[\u0600-\u06FF]', txt):
                        ar_run(tp, txt, size=13)
                    else:
                        add_mixed(tp, txt, size=9)
                        if k == 0:
                            for run in tp.runs: run.bold = True
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(); add_mixed(p, l.strip())
        i += 1
    doc.save(out_path)
    print('OK', out_path)

if __name__ == '__main__':
    render(sys.argv[1], sys.argv[2])
